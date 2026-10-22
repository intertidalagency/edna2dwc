from __future__ import annotations

from dataclasses import dataclass, field
import json
from pathlib import Path
import shutil
import zipfile

from edna2dwc.common.tsv import detect_delimiter, open_text
from edna2dwc.inputs.document_text import UnreadableDocument, document_text
from edna2dwc.inputs.sniff import (
    DELIMITED_SUFFIXES,
    FASTA_SUFFIXES,
    TEXT_BY_CONTENT_SUFFIXES,
    data_suffix,
    looks_delimited,
)

TEXT_PREVIEW_CHARS = 2000
FASTA_PREVIEW_IDS = 5
HEX_PREVIEW_BYTES = 64
XLSX_PREVIEW_ROWS = 5
SAMPLE_BYTES = 65536

_XLSX_EXTS = {".xlsx", ".xlsm"}
_ZIP_EXTS = {".qza", ".qzv", ".zip"}
_DOCUMENT_EXTS = {".md", ".markdown", ".rst", ".pdf", ".docx", ".doc", ".rtf", ".odt"}


@dataclass
class InspectReport:
    path: str
    format: str
    summary: dict = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "path": self.path,
            "format": self.format,
            "summary": self.summary,
            "notes": self.notes,
        }


def inspect_path(path: Path) -> InspectReport:
    suffix = data_suffix(path)
    if suffix in _XLSX_EXTS:
        return _inspect_xlsx(path)
    if suffix in _ZIP_EXTS:
        return _inspect_zip(path)
    if suffix == ".biom":
        return _inspect_biom(path)
    if suffix in DELIMITED_SUFFIXES:
        return _inspect_delimited(path)
    if suffix in TEXT_BY_CONTENT_SUFFIXES:
        return (
            _inspect_delimited(path) if _has_delimited_content(path) else _inspect_document(path)
        )
    if suffix in _DOCUMENT_EXTS:
        return _inspect_document(path)
    if suffix == ".rds":
        return _inspect_rds(path)
    if suffix in FASTA_SUFFIXES:
        return _inspect_fasta(path)
    return _inspect_unrecognized_format(path)


def _inspect_xlsx(path: Path) -> InspectReport:
    import openpyxl  # noqa: PLC0415

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets_summary = [_xlsx_sheet_summary(workbook[name]) for name in workbook.sheetnames]
    finally:
        workbook.close()
    return InspectReport(
        path=str(path),
        format="xlsx",
        summary={"sheets": sheets_summary},
        notes=[
            "Headers may be several rows down, not necessarily row 1.",
            "One sheet per plate/run is common; near-identical sheets are usually one table.",
            "The table may be transposed relative to what its sheet name suggests.",
        ],
    )


def _xlsx_sheet_summary(worksheet) -> dict:
    preview: list[list[str]] = []
    for row in worksheet.iter_rows(values_only=True):
        if len(preview) >= XLSX_PREVIEW_ROWS:
            break
        if any(cell is not None for cell in row):
            preview.append(["" if cell is None else str(cell) for cell in row])
    return {
        "name": worksheet.title,
        "max_row": worksheet.max_row or 0,
        "max_col": worksheet.max_column or 0,
        f"preview_first_{XLSX_PREVIEW_ROWS}_rows": preview,
    }


def _is_qiime_metadata(name: str) -> bool:
    root, _, rest = name.partition("/")
    return bool(root) and rest == "metadata.yaml"


def _inspect_zip(path: Path) -> InspectReport:
    with zipfile.ZipFile(path) as zf:
        namelist = [{"name": info.filename, "size": info.file_size} for info in zf.infolist()]
        if any(_is_qiime_metadata(name) for name in zf.namelist()):
            return _inspect_qza(path, zf, namelist)
    return InspectReport(
        path=str(path),
        format="zip",
        summary={"namelist": namelist},
        notes=[
            "Not a QIIME 2 artifact: it has no <uuid>/metadata.yaml. Unzip it and inspect its files."
        ],
    )


def _inspect_qza(path: Path, zf: zipfile.ZipFile, namelist: list[dict]) -> InspectReport:
    metadata_preview = None
    action_preview = None
    data_previews = {}
    for info in zf.infolist():
        name = info.filename
        if name.endswith("metadata.yaml") and metadata_preview is None:
            metadata_preview = _depth_limited_yaml_preview(
                zf.read(name).decode("utf-8", errors="replace")
            )
        elif name.endswith("action/action.yaml") and action_preview is None:
            action_preview = zf.read(name).decode("utf-8", errors="replace")[:TEXT_PREVIEW_CHARS]
        elif name.endswith((".tsv", ".csv")) and "data/" in name and len(data_previews) < 3:
            data_previews[name] = zf.read(name).decode("utf-8", errors="replace")[
                :TEXT_PREVIEW_CHARS
            ]
    return InspectReport(
        path=str(path),
        format="qza",
        summary={
            "namelist": namelist,
            "metadata_yaml_preview": metadata_preview,
            "first_action_yaml_preview": action_preview,
            "data_file_previews": data_previews,
        },
        notes=[
            (
                "A QIIME 2 artifact is a zip, read with zipfile. Its payload is under "
                "<uuid>/data/: a feature table as feature-table.biom, BIOM 2.x in HDF5 for h5py, "
                "a taxonomy as taxonomy.tsv, and sequences as dna-sequences.fasta."
            )
        ],
    )


def _depth_limited_yaml_preview(text: str, max_depth: int = 2) -> str:
    lines = []
    for line in text.splitlines():
        stripped = line.lstrip(" ")
        indent = len(line) - len(stripped)
        if indent // 2 >= max_depth and stripped.startswith(("-", " ")) is False:
            continue
        if indent // 2 < max_depth:
            lines.append(line)
    return "\n".join(lines)


_HDF5_MAGIC = b"\x89HDF"


def _inspect_biom(path: Path) -> InspectReport:
    with path.open("rb") as f:
        magic = f.read(len(_HDF5_MAGIC))
    if magic == _HDF5_MAGIC:
        return _inspect_biom_hdf5(path)
    with path.open(encoding="utf-8") as f:
        data = json.load(f)
    shape = data.get("shape")
    rows = data.get("rows", [])[:3]
    columns = data.get("columns", [])[:3]
    return InspectReport(
        path=str(path),
        format="biom-json-v1",
        summary={
            "biom_format": data.get("format", "unknown"),
            "shape": shape,
            "first_3_rows": rows,
            "first_3_columns": columns,
        },
    )


def _inspect_biom_hdf5(path: Path) -> InspectReport:
    import h5py  # noqa: PLC0415

    with h5py.File(path, "r") as f:
        obs = f["observation"]
        obs_ids = [_hdf5_decode(v) for v in obs["ids"][:]]
        sample_ids = [_hdf5_decode(v) for v in f["sample"]["ids"][:]]
        rows = [
            {"id": obs_id, "taxonomy": _hdf5_taxonomy(obs, i)}
            for i, obs_id in enumerate(obs_ids[:3])
        ]
        columns = [{"id": sample_id} for sample_id in sample_ids[:3]]
    return InspectReport(
        path=str(path),
        format="biom-hdf5",
        summary={
            "biom_version": "2.x (HDF5)",
            "shape": [len(obs_ids), len(sample_ids)],
            "first_3_rows": rows,
            "first_3_columns": columns,
        },
        notes=["BIOM 2.x in HDF5: counts in `observation/matrix`, read with h5py."],
    )


def _hdf5_decode(value) -> str:
    return value.decode("utf-8") if isinstance(value, bytes) else str(value)


def _hdf5_taxonomy(obs_group, index: int) -> str:
    metadata = obs_group.get("metadata")
    if metadata is None or "taxonomy" not in metadata:
        return ""
    ranks = metadata["taxonomy"][index]
    return ";".join(_hdf5_decode(r) for r in ranks if _hdf5_decode(r))


def _has_delimited_content(path: Path) -> bool:
    try:
        with open_text(path) as f:
            return looks_delimited(f.read(SAMPLE_BYTES))
    except UnicodeDecodeError:
        return False


def _inspect_delimited(path: Path) -> InspectReport:
    with open_text(path) as f:
        sample = f.read(SAMPLE_BYTES)
    delimiter = detect_delimiter(sample)
    lines = sample.splitlines()
    header = lines[0] if lines else ""
    preview_rows = lines[1:6]
    return InspectReport(
        path=str(path),
        format="delimited",
        summary={
            "delimiter": delimiter,
            "header": header.split(delimiter) if header else [],
            "preview_rows": [row.split(delimiter) for row in preview_rows],
        },
    )


def _inspect_rds(path: Path) -> InspectReport:
    rscript = shutil.which("Rscript")
    return InspectReport(
        path=str(path),
        format="rds",
        summary={"rscript_available": rscript is not None},
        notes=[
            (
                "An R object, read with Rscript, so only its presence is checked. A phyloseq "
                "object's otu_table, tax_table and sample_data export to CSV with write.csv."
            )
        ],
    )


DOCUMENT_NOTE = (
    "Documentation: it contains no occurrence records and no reader is written for it. "
    "It is read for dataset metadata (primer sequences, annealing temperature, PCR conditions, "
    "sequencing method, license, creators) when metadata is harvested. `edna2dwc text <file>` "
    "prints all of its text."
)


def _inspect_document(path: Path) -> InspectReport:
    summary: dict = {}
    try:
        summary["text_preview"] = document_text(path, max_chars=TEXT_PREVIEW_CHARS)
    except UnreadableDocument:
        pass
    return InspectReport(path=str(path), format="document", summary=summary, notes=[DOCUMENT_NOTE])


def _unrecognized_format_note(kind: str) -> str:
    return (
        f"Unrecognized {kind}: inspect knows no format with this extension, and the preview is "
        "all it found. The format is identified from context (extension, README, the tool "
        "that produced it)."
    )


def _fasta_records(path: Path) -> dict[str, int]:
    lengths: dict[str, int] = {}
    record_id = None
    with open_text(path) as f:
        for line in f:
            if line.startswith(">"):
                record_id = line[1:].split(maxsplit=1)[0] if line[1:].strip() else ""
                lengths[record_id] = 0
            elif record_id is not None:
                lengths[record_id] += len(line.strip())
    return lengths


def _inspect_fasta(path: Path) -> InspectReport:
    lengths = _fasta_records(path)
    return InspectReport(
        path=str(path),
        format="fasta",
        summary={
            "records": len(lengths),
            "first_ids": list(lengths)[:FASTA_PREVIEW_IDS],
            "length_range": [min(lengths.values(), default=0), max(lengths.values(), default=0)],
        },
        notes=[
            (
                "Sequences keyed by the first word after each '>'. `edna2dwc assemble "
                "--sequences` joins them onto featureID as the sequence column."
            )
        ],
    )


def _inspect_unrecognized_format(path: Path) -> InspectReport:
    try:
        with path.open(encoding="utf-8", errors="strict") as f:
            text = f.read(TEXT_PREVIEW_CHARS)
        return InspectReport(
            path=str(path),
            format="unrecognized-text",
            summary={"text_preview": text},
            notes=[_unrecognized_format_note("format")],
        )
    except UnicodeDecodeError:
        with path.open("rb") as f:
            raw = f.read(HEX_PREVIEW_BYTES)
        return InspectReport(
            path=str(path),
            format="unrecognized-binary",
            summary={"hex_preview": raw.hex()},
            notes=[_unrecognized_format_note("binary format")],
        )
