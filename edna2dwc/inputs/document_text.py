from __future__ import annotations

from collections.abc import Iterator
from datetime import date, datetime
import logging
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

from edna2dwc.common.tsv import open_text
from edna2dwc.inputs.sniff import data_suffix

WORD_NAMESPACE = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
XLSX_SUFFIXES = {".xlsx", ".xlsm"}
BINARY_DOCUMENT_SUFFIXES = {".doc", ".odt"}


class UnreadableDocument(Exception):
    pass


def document_text(path: Path, *, max_chars: int | None = None) -> str:
    chunks: list[str] = []
    length = 0
    for chunk in _chunks(path):
        chunks.append(chunk)
        length += len(chunk)
        if max_chars is not None and length >= max_chars:
            break
    return "".join(chunks)[:max_chars]


def _chunks(path: Path) -> Iterator[str]:
    suffix = data_suffix(path)
    try:
        if suffix == ".pdf":
            yield from _pdf_pages(path)
        elif suffix == ".docx":
            yield from _docx_paragraphs(path)
        elif suffix in XLSX_SUFFIXES:
            yield from _xlsx_sheets(path)
        elif suffix in BINARY_DOCUMENT_SUFFIXES:
            raise UnreadableDocument(f"{path}: save it as .docx or .pdf to read its text")
        else:
            with open_text(path) as f:
                yield f.read()
    except UnreadableDocument:
        raise
    except Exception as error:
        raise UnreadableDocument(f"{path}: {error}") from error


def _pdf_pages(path: Path) -> Iterator[str]:
    from pypdf import PdfReader  # noqa: PLC0415

    logging.getLogger("pypdf").setLevel(logging.ERROR)
    for page in PdfReader(path).pages:
        yield page.extract_text() + "\n"


def _docx_paragraphs(path: Path) -> Iterator[str]:
    with zipfile.ZipFile(path) as archive:
        body = ET.fromstring(archive.read("word/document.xml"))
    for paragraph in body.iter(f"{WORD_NAMESPACE}p"):
        yield "".join(_docx_runs(paragraph)) + "\n"


def _docx_runs(paragraph: ET.Element) -> Iterator[str]:
    for element in paragraph.iter():
        if element.tag == f"{WORD_NAMESPACE}t":
            yield element.text or ""
        elif element.tag == f"{WORD_NAMESPACE}tab":
            yield "\t"
        elif element.tag in (f"{WORD_NAMESPACE}br", f"{WORD_NAMESPACE}cr"):
            yield "\n"


def _xlsx_sheets(path: Path) -> Iterator[str]:
    import openpyxl  # noqa: PLC0415

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            yield f"## {sheet.title}\n"
            for row in sheet.iter_rows(values_only=True):
                yield "\t".join(_cell_text(value) for value in row).rstrip("\t") + "\n"
    finally:
        workbook.close()


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime | date):
        return value.isoformat()
    return " ".join(str(value).split())
