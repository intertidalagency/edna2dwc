from __future__ import annotations

import json
from pathlib import Path
import tempfile
import zipfile

import h5py
import numpy as np
from openpyxl import Workbook

FIXTURES_DIR = Path(__file__).parent / "fixtures" / "inputs"


def _write_xlsx(path: Path) -> None:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Table 1"
    ws1.append(["Station #", "Lat (ddmm.mmm)", "Lon (ddmm.mmm)", "Date"])
    ws1.append(["STA1", 4530.5, -12315.25, "2021-06-01"])
    ws1.append(["STA2", 4531.0, -12316.0, "2021-06-02"])

    ws2 = wb.create_sheet("ASV counts")
    ws2.append(["ASV_ID", "S001", "S002"])
    ws2.append(["asv1", 10, 0])
    ws2.append(["asv2", 5, 20])

    wb.save(path)


def _write_multisheet_xlsx(path: Path) -> None:
    wb = Workbook()
    samples = wb.active
    samples.title = "Samples"
    samples.append(["Sample Name", "Site", "Depth (m)"])
    samples.append(["S001", "North Reef", 12.5])
    samples.append(["S002", "South Reef", 8.0])

    counts = wb.create_sheet("Counts")
    counts.append(["ASV_ID", "S001", "S002"])
    counts.append(["asv1", 10, 0])
    counts.append(["asv2", 5, 20])

    taxonomy = wb.create_sheet("Taxonomy")
    taxonomy.append(["OTU_ID", "Taxon"])
    taxonomy.append(["asv1", "d__Bacteria;p__Cyanobacteria"])
    taxonomy.append(["asv2", "d__Bacteria;p__Firmicutes"])

    wb.save(path)


def _write_biom_json_v1(path: Path) -> None:
    payload = {
        "id": "ingest-fixture",
        "format": "Biological Observation Matrix 1.0.0",
        "format_url": "http://biom-format.org",
        "type": "OTU table",
        "generated_by": "make_input_fixtures.py",
        "date": "2026-08-12T00:00:00",
        "matrix_type": "sparse",
        "matrix_element_type": "int",
        "shape": [2, 2],
        "rows": [
            {"id": "asv1", "metadata": {"taxonomy": ["d__Bacteria", "p__Cyanobacteria"]}},
            {"id": "asv2", "metadata": {"taxonomy": ["d__Bacteria", "p__Firmicutes"]}},
        ],
        "columns": [{"id": "S001", "metadata": None}, {"id": "S002", "metadata": None}],
        "data": [[0, 0, 10], [1, 0, 0], [1, 1, 20]],
    }
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def _write_biom_hdf5(path: Path) -> None:
    with h5py.File(path, "w") as f:
        f.attrs["id"] = "ingest-fixture-hdf5"
        f.attrs["format-url"] = "http://biom-format.org"
        f.attrs["format-version"] = (2, 1)
        f.attrs["generated-by"] = "make_input_fixtures.py"
        f.attrs["creation-date"] = "2026-08-12T00:00:00"
        f.attrs["type"] = "OTU table"
        f.attrs["shape"] = [2, 2]

        obs = f.create_group("observation")
        obs.create_dataset("ids", data=np.array(["asv1", "asv2"], dtype="S"))
        obs_matrix = obs.create_group("matrix")
        obs_matrix.create_dataset("data", data=np.array([10, 3, 20], dtype="f8"))
        obs_matrix.create_dataset("indices", data=np.array([0, 0, 1], dtype="i4"))
        obs_matrix.create_dataset("indptr", data=np.array([0, 1, 3], dtype="i4"))

        obs_metadata = obs.create_group("metadata")
        obs_metadata.create_dataset(
            "taxonomy",
            data=np.array(
                [
                    ["d__Bacteria", "p__Cyanobacteria"],
                    ["d__Bacteria", "p__Firmicutes"],
                ],
                dtype="S32",
            ),
        )

        sample = f.create_group("sample")
        sample.create_dataset("ids", data=np.array(["S001", "S002"], dtype="S"))


def _write_qza(path: Path) -> None:
    uuid = "11111111-1111-1111-1111-111111111111"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(
            f"{uuid}/metadata.yaml",
            "uuid: 11111111-1111-1111-1111-111111111111\n"
            "type: FeatureTable[Frequency]\n"
            "format: BIOMV210DirFmt\n",
        )
        zf.writestr(
            f"{uuid}/data/feature-table.biom",
            "placeholder-biom-bytes",
        )
        zf.writestr(
            f"{uuid}/data/taxonomy.tsv",
            "Feature ID\tTaxon\tConfidence\nasv1\td__Bacteria;p__Cyanobacteria\t0.98\n",
        )
        zf.writestr(
            f"{uuid}/provenance/action/action.yaml",
            "execution:\n"
            "  uuid: 22222222-2222-2222-2222-222222222222\n"
            "action:\n"
            "  type: method\n"
            "  plugin: !ref 'environment:plugins:dada2'\n"
            "  action: denoise-paired\n"
            "  parameters:\n"
            "  - trunc_len_f: 150\n"
            "  - trunc_len_r: 150\n"
            "  - trim_left_f: 0\n"
            "  - chimera_method: consensus\n",
        )
        zf.writestr(
            f"{uuid}/provenance/artifacts/33333333-3333-3333-3333-333333333333/action/action.yaml",
            "execution:\n"
            "  uuid: 44444444-4444-4444-4444-444444444444\n"
            "action:\n"
            "  type: method\n"
            "  plugin: !ref 'environment:plugins:feature-classifier'\n"
            "  action: classify-sklearn\n"
            "  parameters:\n"
            "  - classifier: !ref 'artifacts:55555555-5555-5555-5555-555555555555'\n"
            "  - reads: SILVA 138 99% OTUs full-length sequences\n",
        )
        zf.writestr(
            f"{uuid}/provenance/artifacts/66666666-6666-6666-6666-666666666666/action/action.yaml",
            "execution:\n"
            "  uuid: 77777777-7777-7777-7777-777777777777\n"
            "action:\n"
            "  type: method\n"
            "  plugin: !ref 'environment:plugins:cutadapt'\n"
            "  action: trim-paired\n"
            "  parameters:\n"
            "  - front_f: GTGYCAGCMGCCGCGGTAA\n"
            "  - front_r: GGACTACNVGGGTWTCTAAT\n",
        )


def _write_table_qza(path: Path) -> None:
    uuid = "aaaaaaaa-0000-0000-0000-000000000001"
    with tempfile.TemporaryDirectory() as tmp:
        biom_path = Path(tmp) / "feature-table.biom"
        _write_biom_hdf5(biom_path)
        biom_bytes = biom_path.read_bytes()
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(
            f"{uuid}/metadata.yaml",
            f"uuid: {uuid}\ntype: FeatureTable[Frequency]\nformat: BIOMV210DirFmt\n",
        )
        zf.writestr(f"{uuid}/data/feature-table.biom", biom_bytes)
        zf.writestr(
            f"{uuid}/provenance/action/action.yaml",
            "action:\n"
            "  type: method\n"
            "  plugin: !ref 'environment:plugins:dada2'\n"
            "  action: denoise-paired\n"
            "  parameters:\n"
            "  - trunc_len_f: 150\n",
        )


def _write_delimited(path: Path) -> None:
    path.write_text("sampleID\tfeatureID\tcount\ttaxonomy\ns1\tf1\t10\td__Bacteria\n")


def _write_unrecognized_text(path: Path) -> None:
    path.write_text("A custom LIMS export, tab-separated, no recognized extension.\n")


def _write_unrecognized_binary(path: Path) -> None:
    path.write_bytes(bytes([0xFF, 0xFE, 0x00, 0x01]) + bytes(range(60)))


def write_fixtures(root: Path = FIXTURES_DIR) -> list[Path]:
    root.mkdir(parents=True, exist_ok=True)
    written = []
    for name, writer in [
        ("sample.xlsx", _write_xlsx),
        ("multisheet.xlsx", _write_multisheet_xlsx),
        ("sample_v1.biom", _write_biom_json_v1),
        ("sample_v2.biom", _write_biom_hdf5),
        ("sample.qza", _write_qza),
        ("table.qza", _write_table_qza),
        ("sample.tsv", _write_delimited),
        ("custom_export.lims", _write_unrecognized_text),
        ("custom_export.bin", _write_unrecognized_binary),
    ]:
        path = root / name
        writer(path)
        written.append(path)
    return written


if __name__ == "__main__":
    paths = write_fixtures()
    print(f"Wrote {len(paths)} fixture(s) to {FIXTURES_DIR}")
